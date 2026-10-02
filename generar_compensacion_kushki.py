import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import io
import os
import re
import streamlit as st

# ── Constantes fijas de Kushki ───────────────────────────────────────────
DNI_TERCERO       = "901000330"
TIPO_TERCERO      = "NIT"
CENTRO_COSTO      = "102"

CUENTA_BOLSA      = "138095011"   # débito: total recaudado - comisión - iva
CUENTA_COMISION   = "530515001"   # débito: comisión (1250 por movimiento)
CUENTA_IVA        = "240805002"   # débito: iva de la comisión (19%)
CUENTA_CREDITO    = "110505017"   # crédito: uno a uno por ticket

VALOR_COMISION_UNITARIA = 1250
PORCENTAJE_IVA_COMISION = 0.19

DETALLE_DEBITO = "COMPENSACION PAGO PAGINA KUSHKI"


def _fecha_str(fecha):
    return pd.Timestamp(fecha).strftime("%d-%m-%Y")


def _limpiar_cedula(valor):
    """
    Limpia la cédula para que no quede con el sufijo '.0' que pandas agrega
    cuando una columna numérica se mezcla con celdas vacías (se infiere
    float64 en vez de int). Sin esto, '123456' se convierte en '123456.0'
    y no cruza correctamente contra otra tabla.
    """
    if pd.isna(valor):
        return ""
    s = str(valor).strip()
    if s.endswith(".0"):
        try:
            return str(int(float(s)))
        except Exception:
            return s
    return s


_RE_TICKET_CONCILIACION = re.compile(r"n[uú]mero\s+de\s+ticket\s+(\d+)", re.IGNORECASE)


def _extraer_ticket_descripcion(descripcion):
    """
    Extrae el ticket_number del texto de la columna 'descripcion' del
    archivo de conciliación de canje banco, con el formato típico:
    'Pago realizado a traves de la sucursal virtual con numero de ticket
    <TICKET> y id de transaccion <uuid>'.
    """
    if pd.isna(descripcion):
        return ""
    m = _RE_TICKET_CONCILIACION.search(str(descripcion))
    return m.group(1) if m else ""


def _preparar_indice_conciliacion(archivo_conciliacion):
    """
    Lee el archivo de conciliación de canje banco (el que genera la app en
    Conciliaciones > Cuentas Puentes/Transitorias) y arma un índice por
    ticket_number con la cédula (identificacion) y el tercero registrados
    en el sistema contable para ese movimiento.

    Solo se consideran las filas donde se pudo extraer un ticket_number de
    la columna 'descripcion' — esas son las que vienen de Kushki (usuario
    'Administrador sistema'); el resto del libro de canje banco se ignora.
    """
    df = pd.read_excel(archivo_conciliacion)
    df.columns = [str(c).strip() for c in df.columns]

    requeridas = ["fecha", "descripcion", "identificacion", "tercero", "valor"]
    faltantes = [c for c in requeridas if c not in df.columns]
    if faltantes:
        raise ValueError(
            f"Faltan columnas en el archivo de conciliación: {', '.join(faltantes)}"
        )

    df["_TICKET"] = df["descripcion"].apply(_extraer_ticket_descripcion)
    df = df[df["_TICKET"] != ""].copy()
    if df.empty:
        raise ValueError(
            "El archivo de conciliación no tiene ninguna fila con número de "
            "ticket en la descripción — ¿es el archivo de la cuenta de "
            "canje banco (CANJE BANCOS / Administrador sistema)?"
        )

    df["_IDEN_NORM"]  = df["identificacion"].apply(_limpiar_cedula)
    df["_VALOR_NORM"] = pd.to_numeric(df["valor"], errors="coerce").fillna(0.0)

    duplicados = sorted(df[df.duplicated("_TICKET", keep=False)]["_TICKET"].unique().tolist())

    indice = {}
    for _, row in df.iterrows():
        t = row["_TICKET"]
        if t not in indice:
            indice[t] = {
                "identificacion": row["_IDEN_NORM"],
                "tercero":        str(row.get("tercero", "")).strip(),
                "valor":          row["_VALOR_NORM"],
            }

    return indice, duplicados


def _validar_y_corregir_kushki(df, indice_conciliacion):
    """
    Cruza cada fila del reporte Kushki (por ticket_number) contra el
    índice de la conciliación de canje banco.

    - Si el ticket NO aparece conciliado todavía, esa fila se EXCLUYE de
      la compensación (el pago aún no se aplicó en el sistema).
    - Si aparece, se compara la cédula de Kushki (document_number) contra
      la cédula registrada en la conciliación (identificacion). Si
      coinciden se deja igual; si no coinciden, Kushki trae la cédula del
      medio de pago en vez de la del titular del crédito, así que se usa
      la de la conciliación (la confiable).
    - Si el valor no coincide entre ambos reportes para un ticket ya
      conciliado, se deja pasar pero se reporta como alerta (no bloquea).
    """
    reporte = {"corregidas": [], "sin_cambios": 0, "no_conciliadas": [], "alerta_valor": []}
    filas_validas = []

    for _, row in df.iterrows():
        ticket = str(row["ticket_number"]).strip()
        if ticket.endswith(".0"):
            ticket = ticket[:-2]

        cedula_kushki = _limpiar_cedula(row["document_number"])
        info = indice_conciliacion.get(ticket)

        if info is None:
            reporte["no_conciliadas"].append({
                "ticket_number": ticket,
                "FECHA Kushki":  row["created_dt"].strftime("%d/%m/%Y"),
                "CEDULA Kushki": cedula_kushki,
                "VALOR":         row["approved_transaction_amount"],
            })
            continue

        cedula_final = info["identificacion"] or cedula_kushki
        if cedula_kushki != cedula_final:
            reporte["corregidas"].append({
                "ticket_number":       ticket,
                "CEDULA Kushki (medio de pago)":              cedula_kushki,
                "CEDULA Conciliación (titular del crédito)":  cedula_final,
                "Tercero":             info["tercero"],
            })
        else:
            reporte["sin_cambios"] += 1

        if abs(info["valor"] - float(row["approved_transaction_amount"])) > 1:
            reporte["alerta_valor"].append({
                "ticket_number":       ticket,
                "VALOR Kushki":        row["approved_transaction_amount"],
                "VALOR Conciliación":  info["valor"],
            })

        fila = row.copy()
        fila["document_number"] = cedula_final
        filas_validas.append(fila)

    df_validado = pd.DataFrame(filas_validas) if filas_validas else df.iloc[0:0].copy()
    return df_validado, reporte


def _mostrar_reporte_validacion_kushki(reporte, duplicados):
    st.markdown("---")
    st.markdown("#### 🔎 Validación contra conciliación de canje banco")

    if duplicados:
        st.warning(
            f"⚠️ {len(duplicados)} ticket(s) aparecen más de una vez en el "
            "archivo de conciliación — se usó el primer registro encontrado "
            "para cada uno."
        )

    n_corregidas = len(reporte["corregidas"])
    n_no_concil  = len(reporte["no_conciliadas"])
    n_alerta_val = len(reporte["alerta_valor"])

    if n_corregidas:
        st.error(
            f"🛠️ {n_corregidas} cédula(s) corregida(s): Kushki traía la del "
            "medio de pago; se usó la del titular del crédito registrada en "
            "la conciliación."
        )
        st.dataframe(pd.DataFrame(reporte["corregidas"]), use_container_width=True, hide_index=True)

    if n_no_concil:
        st.warning(
            f"⚠️ {n_no_concil} movimiento(s) de Kushki aún NO aparecen "
            "conciliados en el archivo de canje banco — se excluyeron de "
            "esta compensación. Vuelve a generarla cuando ya estén "
            "conciliados."
        )
        st.dataframe(pd.DataFrame(reporte["no_conciliadas"]), use_container_width=True, hide_index=True)

    if n_alerta_val:
        st.warning(
            f"⚠️ {n_alerta_val} movimiento(s) conciliados pero con un valor "
            "distinto entre Kushki y la conciliación — revisa que sea el "
            "mismo pago antes de subir la compensación."
        )
        st.dataframe(pd.DataFrame(reporte["alerta_valor"]), use_container_width=True, hide_index=True)

    if not n_corregidas and not n_no_concil and not n_alerta_val:
        st.success(f"✅ {reporte['sin_cambios']} movimiento(s) conciliados sin ninguna inconsistencia.")
    elif reporte["sin_cambios"]:
        st.info(f"ℹ️ {reporte['sin_cambios']} movimiento(s) conciliados sin cambios de cédula.")


def crear_compensacion_kushki(archivo_excel, archivo_conciliacion, config=None):
    """
    Genera un archivo de compensación (hoja Items) por cada fecha distinta
    encontrada en la columna 'created' del reporte de Kushki. Solo se
    compensa (no hay aplicación/planos), similar a PSE/Efecty/Record.

    Antes de generar, valida cada movimiento contra el archivo de
    conciliación de canje banco: excluye los que aún no se han aplicado en
    el sistema y corrige la cédula cuando Kushki trae la del medio de pago
    en vez de la del titular del crédito.
    """
    df = pd.read_excel(archivo_excel)
    df.columns = [str(c).strip() for c in df.columns]

    requeridas = ["ticket_number", "created", "approved_transaction_amount", "document_number"]
    faltantes = [c for c in requeridas if c not in df.columns]
    if faltantes:
        raise ValueError(f"Faltan columnas en el archivo: {', '.join(faltantes)}")

    df["created_dt"] = pd.to_datetime(df["created"], errors="coerce", utc=False)
    if df["created_dt"].isna().any():
        n_malas = int(df["created_dt"].isna().sum())
        st.warning(f"⚠️ {n_malas} fila(s) con fecha inválida en 'created' — se excluyeron.")
        df = df[df["created_dt"].notna()].copy()

    if df.empty:
        raise ValueError("No hay filas válidas para procesar.")

    df["approved_transaction_amount"] = pd.to_numeric(
        df["approved_transaction_amount"], errors="coerce"
    ).fillna(0)

    indice_conciliacion, duplicados = _preparar_indice_conciliacion(archivo_conciliacion)
    df, reporte_validacion = _validar_y_corregir_kushki(df, indice_conciliacion)
    _mostrar_reporte_validacion_kushki(reporte_validacion, duplicados)

    if df.empty:
        raise ValueError(
            "Ningún movimiento del reporte Kushki se encontró conciliado en "
            "el archivo de canje banco. No se generó ningún archivo de "
            "compensación."
        )

    df["fecha_grupo"] = df["created_dt"].dt.date

    hora_str = datetime.now().strftime("%d_%m_%Y_%H_%M_%S")
    archivos_generados = []

    for fecha_grupo, df_dia in df.groupby("fecha_grupo"):
        n_movimientos = len(df_dia)
        total_dia      = df_dia["approved_transaction_amount"].sum()
        comision_total = VALOR_COMISION_UNITARIA * n_movimientos
        iva_comision   = round(comision_total * PORCENTAJE_IVA_COMISION, 2)
        valor_bolsa    = round(total_dia - comision_total - iva_comision, 2)

        fecha_dt  = pd.Timestamp(fecha_grupo)
        fecha_txt = _fecha_str(fecha_dt)

        filas_debito = [
            {
                "codigoCentroCosto": CENTRO_COSTO, "dniTercero": DNI_TERCERO,
                "codigoTipoDniTercero": TIPO_TERCERO, "codigoCuenta": CUENTA_BOLSA,
                "valor": valor_bolsa, "factura": fecha_txt, "fechaVencimiento": fecha_dt,
                "codigoImpuesto": None, "valorBaseImpuesto": None, "porcentajeImpuesto": None,
                "detalle": DETALLE_DEBITO,
            },
            {
                "codigoCentroCosto": CENTRO_COSTO, "dniTercero": DNI_TERCERO,
                "codigoTipoDniTercero": TIPO_TERCERO, "codigoCuenta": CUENTA_COMISION,
                "valor": comision_total, "factura": "", "fechaVencimiento": None,
                "codigoImpuesto": None, "valorBaseImpuesto": None, "porcentajeImpuesto": None,
                "detalle": DETALLE_DEBITO,
            },
            {
                "codigoCentroCosto": CENTRO_COSTO, "dniTercero": DNI_TERCERO,
                "codigoTipoDniTercero": TIPO_TERCERO, "codigoCuenta": CUENTA_IVA,
                "valor": iva_comision, "factura": "", "fechaVencimiento": None,
                "codigoImpuesto": "01", "valorBaseImpuesto": comision_total,
                "porcentajeImpuesto": PORCENTAJE_IVA_COMISION,
                "detalle": DETALLE_DEBITO,
            },
        ]

        filas_credito = []
        for _, row in df_dia.iterrows():
            dni_credito = str(row["document_number"]).strip()
            if dni_credito.endswith(".0"):
                dni_credito = dni_credito[:-2]
            ticket = str(row["ticket_number"]).strip()
            if ticket.endswith(".0"):
                ticket = ticket[:-2]
            filas_credito.append({
                "codigoCentroCosto": CENTRO_COSTO, "dniTercero": dni_credito,
                "codigoTipoDniTercero": "CC", "codigoCuenta": CUENTA_CREDITO,
                "valor": -abs(row["approved_transaction_amount"]),
                "factura": "", "fechaVencimiento": None,
                "codigoImpuesto": None, "valorBaseImpuesto": None, "porcentajeImpuesto": None,
                "detalle": f"{DETALLE_DEBITO} {ticket}",
            })

        todas_las_filas = filas_debito + filas_credito
        for idx, fila in enumerate(todas_las_filas, start=1):
            fila["Id"] = idx

        nombre = f"COMPENSACION_KUSHKI_{fecha_dt.strftime('%d_%m_%Y')}_{hora_str}.xlsx"
        buffer = _generar_excel(todas_las_filas)

        ruta_auto = config.get("ruta_compensaciones", "") if config else ""
        if ruta_auto:
            try:
                os.makedirs(ruta_auto, exist_ok=True)
                ruta_completa = os.path.join(ruta_auto, nombre)
                buffer.seek(0)
                with open(ruta_completa, "wb") as f:
                    f.write(buffer.read())
                st.success(f"💾 {nombre}: guardado en {ruta_completa}")
            except Exception as e:
                st.warning(f"⚠️ {nombre}: no se pudo guardar automáticamente — {str(e)}")
            buffer.seek(0)

        archivos_generados.append({"nombre": nombre, "buffer": buffer})

    st.markdown("#### 📥 Descargar compensación(es) generada(s):")
    for arch in archivos_generados:
        arch["buffer"].seek(0)
        st.download_button(
            label=f"⬇️  Descargar {arch['nombre']}",
            data=arch["buffer"],
            file_name=arch["nombre"],
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"dl_kushki_{arch['nombre']}",
        )

    return f"{len(df)} movimiento(s) procesados en {len(archivos_generados)} archivo(s) (uno por fecha)."


def _generar_excel(filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Items"

    encabezados = [
        "Id", "codigoCentroCosto", "dniTercero", "codigoTipoDniTercero",
        "codigoCuenta", "valor", "factura", "fechaVencimiento",
        "codigoImpuesto", "valorBaseImpuesto", "porcentajeImpuesto", "detalle"
    ]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    for col_idx, col_name in enumerate(encabezados, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for fila_idx, fila in enumerate(filas, start=2):
        for col_idx, col_name in enumerate(encabezados, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=fila.get(col_name))
        if fila.get("fechaVencimiento"):
            ws.cell(row=fila_idx, column=8).number_format = "DD/MM/YYYY"

    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
